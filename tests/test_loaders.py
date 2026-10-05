from pathlib import Path

import pytest

from rag_generator.loaders import (
    SUPPORTED_EXTENSIONS,
    EmptyDocumentError,
    UnsupportedFileError,
    load_file,
)


def test_txt_loads_with_source_metadata(tmp_path: Path) -> None:
    f = tmp_path / "a.txt"
    f.write_text("hello world", encoding="utf-8")
    docs = load_file(f, "original.txt")
    assert len(docs) == 1
    assert docs[0].page_content == "hello world"
    assert docs[0].metadata["source"] == "original.txt"


def test_uppercase_md_extension(tmp_path: Path) -> None:
    f = tmp_path / "README.MD"
    f.write_text("# Title\nbody", encoding="utf-8")
    assert load_file(f, "README.MD")[0].metadata["source"] == "README.MD"


@pytest.mark.parametrize("content", ["", "   \n\t\n"])
def test_empty_file_raises(tmp_path: Path, content: str) -> None:
    f = tmp_path / "e.txt"
    f.write_text(content, encoding="utf-8")
    with pytest.raises(EmptyDocumentError, match="No text could be extracted from e.txt"):
        load_file(f, "e.txt")


def test_unknown_extension_raises(tmp_path: Path) -> None:
    f = tmp_path / "a.xyz"
    f.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFileError):
        load_file(f, "a.xyz")


def test_supported_extensions() -> None:
    assert ".txt" in SUPPORTED_EXTENSIONS
    assert ".md" in SUPPORTED_EXTENSIONS


def test_pdf_pages(tmp_path: Path) -> None:
    from fpdf import FPDF

    p = tmp_path / "a.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Hello page one")
    pdf.add_page()
    pdf.cell(text="Page two")
    pdf.output(str(p))
    docs = load_file(p, "a.pdf")
    assert [d.metadata["page"] for d in docs] == [1, 2]
    assert "Hello page one" in docs[0].page_content


def test_docx(tmp_path: Path) -> None:
    from docx import Document as WordDoc

    p = tmp_path / "a.docx"
    d = WordDoc()
    d.add_paragraph("Hello docx")
    d.save(str(p))
    docs = load_file(p, "a.docx")
    assert len(docs) == 1 and "Hello docx" in docs[0].page_content


def test_html_strips_tags(tmp_path: Path) -> None:
    p = tmp_path / "a.html"
    p.write_text("<html><body><h1>Title</h1><p>Body</p></body></html>", encoding="utf-8")
    docs = load_file(p, "a.html")
    assert "<" not in docs[0].page_content
    assert "Title" in docs[0].page_content and "Body" in docs[0].page_content


def test_csv_rows(tmp_path: Path) -> None:
    p = tmp_path / "a.csv"
    p.write_text("name,age\nAnn,30\nBob,\n", encoding="utf-8")
    docs = load_file(p, "a.csv")
    assert [d.metadata["row"] for d in docs] == [1, 2]
    assert all(d.metadata["tabular"] for d in docs)
    assert docs[0].page_content == "name: Ann\nage: 30.0"
    assert docs[1].page_content == "name: Bob"


def test_xlsx_sheets(tmp_path: Path) -> None:
    import pandas as pd

    p = tmp_path / "a.xlsx"
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        pd.DataFrame({"a": [1, 2]}).to_excel(w, sheet_name="S1", index=False)
        pd.DataFrame({"b": ["x"]}).to_excel(w, sheet_name="S2", index=False)
    docs = load_file(p, "a.xlsx")
    assert [(d.metadata["sheet"], d.metadata["row"]) for d in docs] == [("S1", 1), ("S1", 2), ("S2", 1)]
    assert all(d.metadata["tabular"] for d in docs)
